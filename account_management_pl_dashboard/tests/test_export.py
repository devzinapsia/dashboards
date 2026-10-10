import base64
import io

import openpyxl
from freezegun import freeze_time

from odoo.tests import tagged

from .common import PlDashboardCommon


@tagged("post_install", "-at_install")
@freeze_time("2026-03-15")
class TestPlExport(PlDashboardCommon):

    def _sheet(self, **params):
        result = self.Dashboard.get_dashboard_xlsx(**params)
        self.assertTrue(result["filename"].endswith(".xlsx"))
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(result["content"])))
        return workbook.active

    def _row_values(self, sheet, name):
        for row in sheet.iter_rows(values_only=True):
            if row[0] and row[0].strip() == name:
                return row
        self.fail("Row %s not found in the export" % name)

    def test_export_matches_the_grid(self):
        self._entry("2026-01-10", [(self.account_salaries, 100.0, None)])
        self._entry("2026-02-10", [(self.account_sales, -1000.0, self.customer_a), (self.account_rent, 50.0, None)])
        data = self._data()
        sheet = self._sheet(period="fiscal_year")
        salaries = self._row_values(sheet, "Salaries")
        # Name, then the 12 months and the Total, as in the grid.
        self.assertEqual(len([value for value in salaries[1:] if value is not None]), 3 + 1)
        self.assertEqual(salaries[1], 100.0)
        self.assertEqual(salaries[13], self._line_row(data, self.leaf_salaries)["values"]["total"])
        # Groups are indented, percentages exported as ratios.
        self.assertTrue(any(row[0] == "        Salaries" for row in sheet.iter_rows(values_only=True)))
        gross_pct = self._row_values(sheet, "Gross profit %")
        self.assertAlmostEqual(gross_pct[2], 1.0)
        # Profit percentages with two decimals, as in the grid.
        gross_pct_row = next(row for row in sheet.iter_rows() if row[0].value == "Gross profit %")
        self.assertEqual(gross_pct_row[2].number_format, "0.00%")
        self.assertIsNone(gross_pct[1])  # no sales in January: empty, as in the grid

    def test_export_with_budget(self):
        budget = self.env["account.report.budget"].create({
            "name": "Budget 2026", "company_id": self.company.id,
            "item_ids": [(0, 0, {"account_id": self.account_salaries.id, "date": "2026-01-01", "amount": 80.0})],
        })
        self._entry("2026-01-10", [(self.account_salaries, 100.0, None)])
        sheet = self._sheet(period="fiscal_year", budget_id=budget.id)
        salaries = self._row_values(sheet, "Salaries")
        # Actual, budget and deviation for each month.
        self.assertEqual(salaries[1:4], (100.0, 80.0, 0.25))

    def test_month_column_format(self):
        self.assertEqual(self.structure.column_label_format, "mmm_yyyy")
        self.assertEqual(self.Dashboard.get_dashboard_config()["column_label_pattern"], "LLL yyyy")
        self.structure.column_label_format = "yyyy_dash_mm"
        self.assertEqual(self.Dashboard.get_dashboard_config()["column_label_pattern"], "yyyy-MM")
        sheet = self._sheet(period="fiscal_year")
        header = next(row for row in sheet.iter_rows(values_only=True) if row[1] == "2026-01")
        self.assertEqual(header[2], "2026-02")
